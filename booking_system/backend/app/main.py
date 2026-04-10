from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import datetime
from typing import Dict, List
from urllib.parse import quote_plus

from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from sqlalchemy import asc, func, select
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware

from .config import APP_SECRET, BASE_URL
from .database import Base, engine, get_db
from .models import GroupMember, GroupPreference, GroupResult, Meeting, MeetingDate, MeetingGroup, MeetingSlot
from .services.allocator import allocate_fcfs
from .services.emailer import send_mail
from .utils import gen_code, hash_password, parse_datetime_local, split_slots, verify_password

app = FastAPI(title="Teaching Reservation System")
app.add_middleware(SessionMiddleware, secret_key=APP_SECRET, same_site="lax")
app.mount("/static", StaticFiles(directory="app/static"), name="static")
templates = Jinja2Templates(directory="app/templates")


@app.on_event("startup")
def on_startup() -> None:
    Base.metadata.create_all(bind=engine)


def _redirect(url: str, msg: str = "") -> RedirectResponse:
    if msg:
        sep = "&" if "?" in url else "?"
        url = f"{url}{sep}msg={quote_plus(msg)}"
    return RedirectResponse(url=url, status_code=303)


def _teacher_session_key(meeting_code: str) -> str:
    return f"teacher:{meeting_code.strip().upper()}"


def _group_session_key(meeting_code: str, group_code: str) -> str:
    return f"group:{meeting_code.strip().upper()}:{group_code.strip().upper()}"


def _grant_teacher_session(request: Request, meeting_code: str) -> None:
    request.session[_teacher_session_key(meeting_code)] = True


def _grant_group_session(request: Request, meeting_code: str, group_code: str) -> None:
    request.session[_group_session_key(meeting_code, group_code)] = True


def _meeting_by_code(db: Session, meeting_code: str) -> Meeting:
    meeting = db.scalar(select(Meeting).where(Meeting.meeting_code == meeting_code.strip().upper()))
    if not meeting:
        raise HTTPException(status_code=404, detail="Meeting not found")
    return meeting


def _teacher_guard(db: Session, meeting_code: str, password: str) -> Meeting:
    meeting = _meeting_by_code(db, meeting_code)
    if not verify_password(password, meeting.teacher_password):
        raise HTTPException(status_code=403, detail="Invalid teacher password")
    return meeting


def _teacher_session_guard(request: Request, db: Session, meeting_code: str) -> Meeting:
    meeting = _meeting_by_code(db, meeting_code)
    if not request.session.get(_teacher_session_key(meeting_code)):
        raise HTTPException(status_code=403, detail="Teacher session expired")
    return meeting


def _group_guard(db: Session, meeting: Meeting, group_code: str, password: str) -> MeetingGroup:
    group = db.scalar(
        select(MeetingGroup).where(
            MeetingGroup.meeting_id == meeting.id,
            MeetingGroup.group_code == group_code.strip(),
        )
    )
    if not group or not verify_password(password, group.group_password):
        raise HTTPException(status_code=403, detail="Invalid credentials")
    return group


def _group_session_guard(request: Request, db: Session, meeting_code: str, group_code: str) -> tuple[Meeting, MeetingGroup]:
    meeting = _meeting_by_code(db, meeting_code)
    if not request.session.get(_group_session_key(meeting_code, group_code)):
        raise HTTPException(status_code=403, detail="Group session expired")

    group = db.scalar(
        select(MeetingGroup).where(
            MeetingGroup.meeting_id == meeting.id,
            MeetingGroup.group_code == group_code.strip(),
        )
    )
    if not group:
        raise HTTPException(status_code=404, detail="Group not found")
    return meeting, group


def _parse_schedule_lines(raw: str, duration: int):
    slots = []
    for line in raw.splitlines():
        line = line.strip()
        if not line:
            continue
        parts = [x.strip() for x in line.split(",")]
        if len(parts) != 3:
            raise ValueError(f"Invalid schedule line: {line}")
        day, start_hm, end_hm = parts
        slots.extend(split_slots(day, start_hm, end_hm, duration))
    return slots


def _parse_groups_text(raw: str) -> List[dict]:
    rows: List[dict] = []
    for line in raw.splitlines():
        line = line.strip()
        if not line:
            continue

        parts = [p.strip() for p in line.split(",")]
        if len(parts) < 3:
            continue

        group_code = parts[0]
        group_name = parts[1]
        contact_email = parts[2]
        password = parts[3] if len(parts) >= 4 and parts[3] else gen_code("G", 6)
        members_raw = parts[4] if len(parts) >= 5 else ""

        if group_code.lower() in {"group_code", "groupid", "group_id"}:
            continue

        members = [m.strip() for m in members_raw.replace("|", ";").split(";") if m.strip()]
        rows.append(
            {
                "group_code": group_code,
                "group_name": group_name or group_code,
                "contact_email": contact_email,
                "password": password,
                "members": members,
            }
        )
    return rows


def _parse_groups_upload(upload: UploadFile | None) -> List[dict]:
    if not upload or not upload.filename:
        return []

    filename = upload.filename.lower()
    content = upload.file.read()
    out: List[dict] = []

    if filename.endswith(".csv") or filename.endswith(".txt"):
        text = content.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        for row in reader:
            if not row or len(row) < 3:
                continue
            group_code = (row[0] or "").strip()
            group_name = (row[1] or "").strip()
            contact_email = (row[2] or "").strip()
            password = ((row[3] if len(row) >= 4 else "") or "").strip() or gen_code("G", 6)
            members_raw = ((row[4] if len(row) >= 5 else "") or "").strip()
            if group_code.lower() in {"group_code", "groupid", "group_id"}:
                continue
            members = [m.strip() for m in members_raw.replace("|", ";").split(";") if m.strip()]
            out.append(
                {
                    "group_code": group_code,
                    "group_name": group_name or group_code,
                    "contact_email": contact_email,
                    "password": password,
                    "members": members,
                }
            )
        return out

    if filename.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=1, values_only=True):
            group_code = str(row[0]).strip() if len(row) >= 1 and row[0] is not None else ""
            group_name = str(row[1]).strip() if len(row) >= 2 and row[1] is not None else ""
            contact_email = str(row[2]).strip() if len(row) >= 3 and row[2] is not None else ""
            password = str(row[3]).strip() if len(row) >= 4 and row[3] is not None else gen_code("G", 6)
            members_raw = str(row[4]).strip() if len(row) >= 5 and row[4] is not None else ""
            if group_code.lower() in {"group_code", "groupid", "group_id"}:
                continue
            members = [m.strip() for m in members_raw.replace("|", ";").split(";") if m.strip()]
            out.append(
                {
                    "group_code": group_code,
                    "group_name": group_name or group_code,
                    "contact_email": contact_email,
                    "password": password,
                    "members": members,
                }
            )
        return out

    return []


def _validate_preference_constraints(db: Session, meeting: Meeting, slot_ids: List[int]) -> tuple[List[int], str]:
    if not slot_ids:
        return [], ""

    dedup: List[int] = []
    seen = set()
    for sid in slot_ids:
        if sid not in seen:
            dedup.append(sid)
            seen.add(sid)

    max_total = meeting.date_choice_num * meeting.slot_choice_num
    dedup = dedup[:max_total]

    valid_slots = db.scalars(
        select(MeetingSlot)
        .where(
            MeetingSlot.id.in_(dedup),
            MeetingSlot.meeting_id == meeting.id,
            MeetingSlot.scheduled.is_(False),
        )
        .order_by(MeetingSlot.day, MeetingSlot.start_time)
    ).all()

    slot_map = {s.id: s for s in valid_slots}
    filtered = [sid for sid in dedup if sid in slot_map]

    day_counts: Dict[str, int] = defaultdict(int)
    used_days = set()
    for sid in filtered:
        day_key = str(slot_map[sid].day)
        day_counts[day_key] += 1
        used_days.add(day_key)

    if len(used_days) > meeting.date_choice_num:
        return [], f"You selected {len(used_days)} days, max allowed is {meeting.date_choice_num}."

    for day_key, count in day_counts.items():
        if count > meeting.slot_choice_num:
            return [], f"Day {day_key} has {count} slots, max allowed per day is {meeting.slot_choice_num}."

    return filtered, ""


def _deliver_mail_batch(messages: list[tuple[str, str, str]]) -> dict[str, int]:
    sent = 0
    failed = 0

    for to_addr, subject, body in messages:
        if not to_addr:
            continue
        try:
            send_mail(to_addr, subject, body)
            sent += 1
        except Exception as exc:
            failed += 1
            print(f"[MAIL_ERROR] to={to_addr} subject={subject} error={exc}")

    return {"sent": sent, "failed": failed}


def _build_creation_mail_jobs(
    meeting: Meeting,
    raw_teacher_password: str,
    groups: list[dict],
    slot_count: int,
) -> list[tuple[str, str, str]]:
    jobs = [
        (
            meeting.teacher_email,
            f"[Reservation] Meeting {meeting.title} created",
            (
                f"Your teaching reservation meeting has been created.\n\n"
                f"Meeting title: {meeting.title}\n"
                f"Meeting code: {meeting.meeting_code}\n"
                f"Teacher password: {raw_teacher_password}\n"
                f"Deadline: {meeting.deadline}\n"
                f"Total generated slots: {slot_count}\n"
                f"Dashboard URL: {BASE_URL}/teacher/login\n"
            ),
        )
    ]

    for group in groups:
        jobs.append(
            (
                group["contact_email"],
                f"[Reservation] Group credentials for {meeting.title}",
                (
                    f"Your group can now submit reservation preferences.\n\n"
                    f"Meeting title: {meeting.title}\n"
                    f"Meeting code: {meeting.meeting_code}\n"
                    f"Group code: {group['group_code']}\n"
                    f"Group password: {group['password']}\n"
                    f"Deadline: {meeting.deadline}\n"
                    f"Group login URL: {BASE_URL}/group/login\n"
                    f"Please keep these credentials confidential.\n"
                ),
            )
        )

    return jobs


def _build_result_mail_jobs(db: Session, meeting: Meeting, round_index: int, allocated_group_ids: list[int]) -> list[tuple[str, str, str]]:
    jobs: list[tuple[str, str, str]] = []

    allocated_rows = db.execute(
        select(
            MeetingGroup.group_code,
            MeetingGroup.group_name,
            MeetingGroup.contact_email,
            MeetingSlot.day,
            MeetingSlot.start_time,
            MeetingSlot.end_time,
        )
        .join(GroupResult, GroupResult.meeting_group_id == MeetingGroup.id)
        .join(MeetingSlot, MeetingSlot.id == GroupResult.slot_id)
        .where(
            GroupResult.meeting_id == meeting.id,
            GroupResult.round_index == round_index,
            GroupResult.meeting_group_id.in_(allocated_group_ids or [-1]),
        )
        .order_by(MeetingGroup.group_code)
    ).all()

    for row in allocated_rows:
        jobs.append(
            (
                row[2],
                f"[Reservation] Allocation result for {meeting.title}",
                (
                    f"Your group has been assigned a demo slot.\n\n"
                    f"Meeting title: {meeting.title}\n"
                    f"Meeting code: {meeting.meeting_code}\n"
                    f"Group code: {row[0]}\n"
                    f"Assigned date: {row[3]}\n"
                    f"Assigned time: {row[4]} - {row[5]}\n"
                    f"Allocation round: {round_index}\n"
                    f"Result portal: {BASE_URL}/group/login\n"
                ),
            )
        )

    pending_groups = db.scalars(
        select(MeetingGroup)
        .where(MeetingGroup.meeting_id == meeting.id, MeetingGroup.scheduled.is_(False))
        .order_by(MeetingGroup.group_code)
    ).all()

    for group in pending_groups:
        jobs.append(
            (
                group.contact_email,
                f"[Reservation] Round {round_index} result pending for {meeting.title}",
                (
                    f"No slot was assigned to your group in round {round_index}.\n\n"
                    f"Meeting title: {meeting.title}\n"
                    f"Meeting code: {meeting.meeting_code}\n"
                    f"Group code: {group.group_code}\n"
                    f"You will receive another email if a next round is opened.\n"
                    f"Group login URL: {BASE_URL}/group/login\n"
                ),
            )
        )

    jobs.append(
        (
            meeting.teacher_email,
            f"[Reservation] Allocation summary for {meeting.title}",
            (
                f"Round {round_index} allocation has finished.\n\n"
                f"Meeting code: {meeting.meeting_code}\n"
                f"Newly allocated groups: {len(allocated_group_ids)}\n"
                f"Remaining unscheduled groups: {len(pending_groups)}\n"
                f"Dashboard URL: {BASE_URL}/teacher/login\n"
            ),
        )
    )

    return jobs


def _build_next_round_mail_jobs(meeting: Meeting, unscheduled_groups: list[MeetingGroup], extra_schedule_lines: str) -> list[tuple[str, str, str]]:
    jobs: list[tuple[str, str, str]] = []
    extra_schedule_summary = extra_schedule_lines.strip() or "No additional slots were added."

    for group in unscheduled_groups:
        jobs.append(
            (
                group.contact_email,
                f"[Reservation] Round {meeting.round_index} is open for {meeting.title}",
                (
                    f"A new reservation round has been opened for your group.\n\n"
                    f"Meeting title: {meeting.title}\n"
                    f"Meeting code: {meeting.meeting_code}\n"
                    f"Group code: {group.group_code}\n"
                    f"New deadline: {meeting.deadline}\n"
                    f"Round index: {meeting.round_index}\n"
                    f"Additional schedule:\n{extra_schedule_summary}\n\n"
                    f"Group login URL: {BASE_URL}/group/login\n"
                ),
            )
        )

    jobs.append(
        (
            meeting.teacher_email,
            f"[Reservation] Round {meeting.round_index} opened for {meeting.title}",
            (
                f"Next round has been opened successfully.\n\n"
                f"Meeting code: {meeting.meeting_code}\n"
                f"New deadline: {meeting.deadline}\n"
                f"Remaining unscheduled groups: {len(unscheduled_groups)}\n"
                f"Dashboard URL: {BASE_URL}/teacher/login\n"
            ),
        )
    )
    return jobs


@app.get("/")
def home(request: Request):
    return templates.TemplateResponse("home.html", {"request": request})


@app.get("/teacher/login")
def teacher_login_page(request: Request, msg: str = Query("")):
    return templates.TemplateResponse("teacher_login.html", {"request": request, "msg": msg})


@app.post("/teacher/login")
def teacher_login(
    request: Request,
    meeting_code: str = Form(...),
    password: str = Form(...),
    db: Session = Depends(get_db),
):
    try:
        meeting = _teacher_guard(db, meeting_code, password)
    except HTTPException:
        return _redirect("/teacher/login", "Invalid meeting code or password")

    _grant_teacher_session(request, meeting.meeting_code)
    return _redirect(f"/teacher/{meeting.meeting_code}/dashboard")


@app.get("/teacher/create")
def teacher_create_page(request: Request):
    return templates.TemplateResponse(
        "teacher_create.html",
        {
            "request": request,
            "tomorrow": datetime.now().replace(hour=23, minute=59).strftime("%Y-%m-%dT%H:%M"),
        },
    )


@app.post("/teacher/create")
def teacher_create(
    request: Request,
    title: str = Form(...),
    teacher_email: str = Form(...),
    duration_minutes: int = Form(...),
    deadline: str = Form(...),
    date_choice_num: int = Form(2),
    slot_choice_num: int = Form(3),
    schedule_lines: str = Form(...),
    groups_lines: str = Form(""),
    groups_file: UploadFile | None = File(default=None),
    db: Session = Depends(get_db),
):
    try:
        deadline_dt = parse_datetime_local(deadline)
    except ValueError:
        return _redirect("/teacher/create", "Invalid deadline format")

    if duration_minutes < 5:
        return _redirect("/teacher/create", "Duration should be >= 5 minutes")

    try:
        slots = _parse_schedule_lines(schedule_lines, duration_minutes)
    except ValueError as e:
        return _redirect("/teacher/create", str(e))

    if not slots:
        return _redirect("/teacher/create", "No valid slots generated from schedule")

    groups = _parse_groups_text(groups_lines)
    groups.extend(_parse_groups_upload(groups_file))

    unique_groups: Dict[str, dict] = {}
    for g in groups:
        if g["group_code"] and g["contact_email"] and g["group_code"] not in unique_groups:
            unique_groups[g["group_code"]] = g

    if not unique_groups:
        return _redirect("/teacher/create", "No valid groups provided")

    raw_teacher_password = gen_code("T", 8)
    meeting = Meeting(
        meeting_code=gen_code("M", 8),
        title=title.strip(),
        teacher_email=teacher_email.strip(),
        teacher_password=hash_password(raw_teacher_password),
        duration_minutes=duration_minutes,
        deadline=deadline_dt,
        date_choice_num=max(1, date_choice_num),
        slot_choice_num=max(1, slot_choice_num),
    )
    db.add(meeting)
    db.flush()

    unique_days = sorted({d for d, _, _ in slots})
    for day in unique_days:
        db.add(MeetingDate(meeting_id=meeting.id, day=day))

    for day, start_hm, end_hm in slots:
        db.add(MeetingSlot(meeting_id=meeting.id, day=day, start_time=start_hm, end_time=end_hm, round_index=meeting.round_index))

    group_payloads = list(unique_groups.values())
    for g in group_payloads:
        mg = MeetingGroup(
            meeting_id=meeting.id,
            group_code=g["group_code"],
            group_name=g["group_name"],
            contact_email=g["contact_email"],
            group_password=hash_password(g["password"]),
        )
        db.add(mg)
        db.flush()

        for member in g["members"]:
            db.add(GroupMember(meeting_group_id=mg.id, member_id=member, member_name=""))

    db.commit()

    _grant_teacher_session(request, meeting.meeting_code)
    mail_summary = _deliver_mail_batch(
        _build_creation_mail_jobs(meeting, raw_teacher_password, group_payloads, len(slots))
    )
    return templates.TemplateResponse(
        "teacher_created.html",
        {
            "request": request,
            "meeting_code": meeting.meeting_code,
            "password": raw_teacher_password,
            "msg": "Meeting created successfully",
            "mail_summary": mail_summary,
            "group_count": len(group_payloads),
            "slot_count": len(slots),
        },
    )


@app.get("/teacher/created")
def teacher_created_page():
    return _redirect("/teacher/create", "Credentials are shown immediately after meeting creation")


@app.get("/teacher/{meeting_code}/dashboard")
def teacher_dashboard(
    request: Request,
    meeting_code: str,
    msg: str = Query(""),
    db: Session = Depends(get_db),
):
    try:
        meeting = _teacher_session_guard(request, db, meeting_code)
    except HTTPException:
        return _redirect("/teacher/login", "Please sign in to continue")

    slots = db.scalars(select(MeetingSlot).where(MeetingSlot.meeting_id == meeting.id).order_by(MeetingSlot.day, MeetingSlot.start_time)).all()
    groups = db.scalars(select(MeetingGroup).where(MeetingGroup.meeting_id == meeting.id).order_by(MeetingGroup.group_code)).all()

    member_counts = dict(
        db.execute(
            select(GroupMember.meeting_group_id, func.count(GroupMember.id)).group_by(GroupMember.meeting_group_id)
        ).all()
    )

    result_rows = db.execute(
        select(MeetingGroup.group_code, MeetingGroup.group_name, MeetingSlot.day, MeetingSlot.start_time, MeetingSlot.end_time, GroupResult.round_index)
        .join(GroupResult, GroupResult.meeting_group_id == MeetingGroup.id)
        .join(MeetingSlot, MeetingSlot.id == GroupResult.slot_id)
        .where(GroupResult.meeting_id == meeting.id)
        .order_by(MeetingGroup.group_code)
    ).all()

    total_groups = db.scalar(select(func.count(MeetingGroup.id)).where(MeetingGroup.meeting_id == meeting.id)) or 0
    scheduled_groups = db.scalar(select(func.count(MeetingGroup.id)).where(MeetingGroup.meeting_id == meeting.id, MeetingGroup.scheduled.is_(True))) or 0
    submitted_current = (
        db.scalar(select(func.count(func.distinct(GroupPreference.meeting_group_id))).where(GroupPreference.meeting_id == meeting.id, GroupPreference.round_index == meeting.round_index))
        or 0
    )
    free_slots = db.scalar(select(func.count(MeetingSlot.id)).where(MeetingSlot.meeting_id == meeting.id, MeetingSlot.scheduled.is_(False))) or 0

    return templates.TemplateResponse(
        "teacher_dashboard.html",
        {
            "request": request,
            "meeting": meeting,
            "slots": slots,
            "groups": groups,
            "member_counts": member_counts,
            "result_rows": result_rows,
            "msg": msg,
            "stats": {
                "total_groups": total_groups,
                "scheduled_groups": scheduled_groups,
                "submitted_current": submitted_current,
                "free_slots": free_slots,
                "unscheduled_groups": total_groups - scheduled_groups,
            },
        },
    )


@app.post("/teacher/{meeting_code}/allocate")
def teacher_allocate(request: Request, meeting_code: str, db: Session = Depends(get_db)):
    try:
        meeting = _teacher_session_guard(request, db, meeting_code)
    except HTTPException:
        return _redirect("/teacher/login", "Please sign in to continue")

    summary = allocate_fcfs(db, meeting)
    mail_summary = _deliver_mail_batch(
        _build_result_mail_jobs(db, meeting, summary["round_index"], summary["allocated_group_ids"])
    )

    msg = (
        f"Round {summary['round_index']} finished. Newly allocated={summary['allocated_count']}, "
        f"unscheduled={summary['unscheduled_count']}, emails sent={mail_summary['sent']}, failed={mail_summary['failed']}"
    )
    return _redirect(f"/teacher/{meeting.meeting_code}/dashboard", msg)


@app.post("/teacher/{meeting_code}/next-round")
def next_round(
    request: Request,
    meeting_code: str,
    new_deadline: str = Form(""),
    extra_schedule_lines: str = Form(""),
    db: Session = Depends(get_db),
):
    try:
        meeting = _teacher_session_guard(request, db, meeting_code)
    except HTTPException:
        return _redirect("/teacher/login", "Please sign in to continue")

    if new_deadline.strip():
        try:
            meeting.deadline = parse_datetime_local(new_deadline.strip())
        except ValueError:
            return _redirect(f"/teacher/{meeting.meeting_code}/dashboard", "Invalid new deadline")

    meeting.round_index += 1

    if extra_schedule_lines.strip():
        try:
            slots = _parse_schedule_lines(extra_schedule_lines, meeting.duration_minutes)
        except ValueError as e:
            return _redirect(f"/teacher/{meeting.meeting_code}/dashboard", str(e))

        unique_days = sorted({d for d, _, _ in slots})
        for day in unique_days:
            exists = db.scalar(select(MeetingDate).where(MeetingDate.meeting_id == meeting.id, MeetingDate.day == day))
            if not exists:
                db.add(MeetingDate(meeting_id=meeting.id, day=day))

        for day, start_hm, end_hm in slots:
            db.add(MeetingSlot(meeting_id=meeting.id, day=day, start_time=start_hm, end_time=end_hm, round_index=meeting.round_index))

    db.commit()

    unscheduled_groups = db.scalars(
        select(MeetingGroup)
        .where(MeetingGroup.meeting_id == meeting.id, MeetingGroup.scheduled.is_(False))
        .order_by(MeetingGroup.group_code)
    ).all()
    mail_summary = _deliver_mail_batch(
        _build_next_round_mail_jobs(meeting, unscheduled_groups, extra_schedule_lines)
    )

    msg = (
        f"Switched to round {meeting.round_index}. "
        f"Emails sent={mail_summary['sent']}, failed={mail_summary['failed']}"
    )
    return _redirect(f"/teacher/{meeting.meeting_code}/dashboard", msg)


@app.get("/group/login")
def group_login_page(request: Request, msg: str = Query("")):
    return templates.TemplateResponse("group_login.html", {"request": request, "msg": msg})


@app.post("/group/login")
def group_login(
    request: Request,
    meeting_code: str = Form(...),
    group_code: str = Form(...),
    password: str = Form(...),
    db: Session = Depends(get_db),
):
    try:
        meeting = _meeting_by_code(db, meeting_code)
        grp = _group_guard(db, meeting, group_code, password)
    except HTTPException:
        return _redirect("/group/login", "Invalid credentials")

    _grant_group_session(request, meeting.meeting_code, grp.group_code)
    return _redirect(f"/group/{meeting.meeting_code}/{grp.group_code}/preference")


@app.get("/group/{meeting_code}/{group_code}/preference")
def group_preference_page(
    request: Request,
    meeting_code: str,
    group_code: str,
    msg: str = Query(""),
    db: Session = Depends(get_db),
):
    try:
        meeting, grp = _group_session_guard(request, db, meeting_code, group_code)
    except HTTPException:
        return _redirect("/group/login", "Please sign in again")

    now = datetime.now()
    editable = now <= meeting.deadline and not grp.scheduled

    slots = db.scalars(
        select(MeetingSlot)
        .where(MeetingSlot.meeting_id == meeting.id, MeetingSlot.scheduled.is_(False))
        .order_by(MeetingSlot.day, MeetingSlot.start_time)
    ).all()

    old_prefs = db.scalars(
        select(GroupPreference)
        .where(
            GroupPreference.meeting_id == meeting.id,
            GroupPreference.meeting_group_id == grp.id,
            GroupPreference.round_index == meeting.round_index,
        )
        .order_by(GroupPreference.priority)
    ).all()

    selected_slot_ids = [p.slot_id for p in old_prefs]
    slot_by_id = {slot.id: slot for slot in slots}
    old_pref_rows = [slot_by_id[sid] for sid in selected_slot_ids if sid in slot_by_id]

    grouped_slots: Dict[str, list[MeetingSlot]] = defaultdict(list)
    for slot in slots:
        grouped_slots[str(slot.day)].append(slot)

    members = db.scalars(select(GroupMember).where(GroupMember.meeting_group_id == grp.id).order_by(GroupMember.member_id)).all()

    return templates.TemplateResponse(
        "group_preference.html",
        {
            "request": request,
            "meeting": meeting,
            "group": grp,
            "members": members,
            "editable": editable,
            "selected_slot_ids": selected_slot_ids,
            "old_pref_rows": old_pref_rows,
            "grouped_slots": dict(grouped_slots),
            "msg": msg,
            "now": now,
            "max_days": meeting.date_choice_num,
            "max_per_day": meeting.slot_choice_num,
            "max_total": meeting.date_choice_num * meeting.slot_choice_num,
        },
    )


@app.post("/group/{meeting_code}/{group_code}/preference")
def group_submit_preference(
    request: Request,
    meeting_code: str,
    group_code: str,
    slot_ids: str = Form(""),
    db: Session = Depends(get_db),
):
    try:
        meeting, grp = _group_session_guard(request, db, meeting_code, group_code)
    except HTTPException:
        return _redirect("/group/login", "Please sign in again")

    if datetime.now() > meeting.deadline:
        return _redirect(f"/group/{meeting.meeting_code}/{grp.group_code}/preference", "Deadline passed")

    if grp.scheduled:
        return _redirect(f"/group/{meeting.meeting_code}/{grp.group_code}/preference", "Group already scheduled")

    parsed = []
    for raw in slot_ids.split(","):
        raw = raw.strip()
        if raw.isdigit():
            parsed.append(int(raw))

    filtered, err = _validate_preference_constraints(db, meeting, parsed)
    if err:
        return _redirect(f"/group/{meeting.meeting_code}/{grp.group_code}/preference", err)

    db.query(GroupPreference).filter(
        GroupPreference.meeting_id == meeting.id,
        GroupPreference.meeting_group_id == grp.id,
        GroupPreference.round_index == meeting.round_index,
    ).delete(synchronize_session=False)

    for idx, sid in enumerate(filtered, start=1):
        db.add(
            GroupPreference(
                meeting_id=meeting.id,
                meeting_group_id=grp.id,
                slot_id=sid,
                priority=idx,
                round_index=meeting.round_index,
            )
        )

    grp.last_submitted_at = datetime.now()
    db.commit()

    return _redirect(f"/group/{meeting.meeting_code}/{grp.group_code}/preference", "Preference saved")


@app.get("/group/{meeting_code}/{group_code}/result")
def group_result(request: Request, meeting_code: str, group_code: str, db: Session = Depends(get_db)):
    try:
        meeting, grp = _group_session_guard(request, db, meeting_code, group_code)
    except HTTPException:
        return _redirect("/group/login", "Please sign in again")

    row = db.execute(
        select(MeetingSlot.day, MeetingSlot.start_time, MeetingSlot.end_time, GroupResult.round_index)
        .join(GroupResult, GroupResult.slot_id == MeetingSlot.id)
        .where(GroupResult.meeting_id == meeting.id, GroupResult.meeting_group_id == grp.id)
        .order_by(asc(GroupResult.assigned_at))
    ).first()

    members = db.scalars(select(GroupMember).where(GroupMember.meeting_group_id == grp.id).order_by(GroupMember.member_id)).all()

    return templates.TemplateResponse(
        "group_result.html",
        {
            "request": request,
            "meeting": meeting,
            "group": grp,
            "members": members,
            "result": row,
        },
    )


@app.get("/student/login")
def student_deprecated():
    return _redirect("/group/login", "System upgraded to group reservation login")
